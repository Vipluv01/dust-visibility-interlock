"""Render the report to HTML (for the PDF export) and to .docx, from one source.

    .venv/bin/python writeup/build.py
"""
import html as _html
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "writeup"))

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import content as C

OUT_HTML = ROOT / "writeup" / "report.html"
OUT_DOCX = ROOT / "Dust-Visibility-Interlock.docx"


# --------------------------------------------------------------------------
# the governor excerpt the brief asks to be pasted into Part 5
# --------------------------------------------------------------------------
def make_excerpt():
    src = (ROOT / "governor.py").read_text().splitlines()
    want = []
    grab = False
    for ln in src:
        if ln.startswith("def safe_speed_mph"):
            grab = True
        if ln.startswith("@dataclass") and want:
            grab = False
        if ln.startswith("def plume_detected") or ln.startswith("class VisibilityGovernor"):
            grab = True
        if ln.startswith("def brief_interlock"):
            grab = False
        if grab:
            want.append(ln)
    text = "\n".join(want).strip()
    (ROOT / "writeup" / "governor_excerpt.txt").write_text(text)
    return text


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------
CSS = """
:root{--ink:#17222e;--ink2:#3c4personally;}
"""

HTML_CSS = """
<style>
  :root { --ink:#17222e; --ink-2:#44525f; --ink-3:#73818d; --rule:#c9d3dc;
          --accent:#1f4e79; --warm:#9c3d1e; --paper:#ffffff; --tint:#f4f7fa; }
  * { box-sizing: border-box; }
  body { margin:0; background:var(--paper); color:var(--ink);
         font-family: "Charter","Georgia",'Times New Roman',serif;
         font-size:11.2pt; line-height:1.62; }
  .wrap { max-width: 46rem; margin:0 auto; padding: 3.2rem 1.6rem 5rem; }
  h1 { font-size: 2.05rem; line-height:1.18; margin:0 0 .5rem; letter-spacing:-.012em; }
  .sub { font-size:1.02rem; color:var(--ink-2); margin:0 0 1.5rem; line-height:1.45; }
  .cover-meta { font-family: ui-monospace,"SF Mono",Menlo,monospace; font-size:.72rem;
                letter-spacing:.06em; text-transform:uppercase; color:var(--ink-3);
                border-top:1px solid var(--rule); border-bottom:1px solid var(--rule);
                padding:.75rem 0; margin:0 0 2.4rem; }
  .cover-meta strong { display:block; color:var(--ink); font-size:.86rem;
                       letter-spacing:.03em; text-transform:none; margin-bottom:.3rem; }
  h2 { font-size:1.32rem; margin:2.6rem 0 .7rem; padding-top:.9rem;
       border-top:2px solid var(--ink); letter-spacing:-.008em; }
  h3 { font-size:1.04rem; margin:1.8rem 0 .5rem; color:var(--accent); }
  p { margin:0 0 .85rem; }
  code { font-family: ui-monospace,"SF Mono",Menlo,monospace; font-size:.86em;
         background:var(--tint); padding:.08em .32em; border-radius:3px; }
  ul { margin:0 0 1rem; padding-left:1.15rem; }
  li { margin-bottom:.42rem; }
  figure { margin:1.5rem 0; }
  figure img { width:100%; height:auto; display:block; }
  figcaption { font-size:.84rem; color:var(--ink-2); margin-top:.55rem;
               line-height:1.45; }
  table { width:100%; border-collapse:collapse; margin:.3rem 0 .4rem;
          font-size:.80rem; font-variant-numeric: tabular-nums; }
  caption { caption-side:top; text-align:left; font-size:.84rem; color:var(--ink-2);
            margin-bottom:.45rem; line-height:1.45; }
  thead th { text-align:left; font-weight:600; font-size:.74rem;
             text-transform:uppercase; letter-spacing:.04em; color:var(--ink-2);
             border-top:1.4px solid var(--ink); border-bottom:1px solid var(--ink);
             padding:.46rem .5rem; vertical-align:bottom; }
  tbody td { padding:.4rem .5rem; border:none; vertical-align:top; }
  tbody tr:last-child td { border-bottom:1.4px solid var(--ink); }
  td:not(:first-child), th:not(:first-child) { text-align:right; }
  .note { border-left:3px solid var(--warm); background:#fdf5f1; padding:.8rem 1rem;
          margin:1.1rem 0; font-size:.94rem; }
  pre { background:var(--tint); border:1px solid var(--rule); border-radius:4px;
        padding:.85rem 1rem; overflow-x:auto; font-size:.70rem; line-height:1.5;
        font-family: ui-monospace,"SF Mono",Menlo,monospace; }
  @page { size:A4; margin:16mm 15mm 15mm; }
  @media print {
    body { font-size:9.9pt; line-height:1.48; }
    .wrap { max-width:none; padding:0; }
    h1 { font-size:21pt; }
    h2 { font-size:13pt; margin-top:1.5rem; break-after:avoid; }
    h3 { font-size:11pt; break-after:avoid; }
    table { font-size:7.4pt; }
    thead th { font-size:6.5pt; }
    pre { font-size:6.4pt; white-space:pre-wrap; word-break:break-word; }
    figure, table, .note { break-inside:avoid; }
    tr, td, th { break-inside:avoid; }
    caption { font-size:7.6pt; }
    figcaption { font-size:7.6pt; }
    h2 { break-before:auto; }
  }
</style>
"""


def inline(t):
    """**bold** -> <strong>, leave existing tags alone."""
    return re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)


def render_html(bl):
    out = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'>",
           f"<meta name='viewport' content='width=device-width,initial-scale=1'>",
           f"<title>{_html.escape(C.TITLE)}</title>{HTML_CSS}</head><body><div class='wrap'>"]
    out.append(f"<h1>{_html.escape(C.TITLE)}</h1>")
    out.append(f"<p class='sub'>{_html.escape(C.SUBTITLE)}</p>")
    out.append(f"<p class='cover-meta'><strong>{C.AUTHOR} &middot; Roll no. {C.ROLL}</strong>"
               f"{C.AFFIL} &middot; Capstone Project 10</p>")
    for blk in bl:
        k = blk[0]
        if k == "h2":
            out.append(f"<h2>{blk[1]}</h2>")
        elif k == "h3":
            out.append(f"<h3>{blk[1]}</h3>")
        elif k == "p":
            out.append(f"<p>{inline(blk[1])}</p>")
        elif k == "note":
            out.append(f"<div class='note'>{inline(blk[1])}</div>")
        elif k == "bullets":
            items = "".join(f"<li>{inline(i)}</li>" for i in blk[1])
            out.append(f"<ul>{items}</ul>")
        elif k == "pre":
            if blk[1].strip():
                out.append(f"<pre>{_html.escape(blk[1])}</pre>")
        elif k == "fig":
            out.append(f"<figure><img src='../{blk[1]}' alt=''>"
                       f"<figcaption>{blk[2]}</figcaption></figure>")
        elif k == "table":
            _, cap, hdr, rows = blk
            th = "".join(f"<th>{h}</th>" for h in hdr)
            tr = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>"
                         for r in rows)
            out.append(f"<table><caption>{cap}</caption>"
                       f"<thead><tr>{th}</tr></thead><tbody>{tr}</tbody></table>")
    out.append("</div></body></html>")
    OUT_HTML.write_text("\n".join(out))
    return OUT_HTML


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def strip_tags(t):
    t = t.replace("&ldquo;", "“").replace("&rdquo;", "”")
    t = t.replace("&middot;", "·").replace("&times;", "×")
    t = re.sub(r"</?(em|strong|code)>", "", t)
    t = re.sub(r"<sub>(.*?)</sub>", r"\1", t)
    t = re.sub(r"<[^>]+>", "", t)
    return _html.unescape(t)


def rich_runs(par, text):
    """Emit runs, honouring <strong>/<em>/<code> and **bold**."""
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    text = text.replace("&ldquo;", "“").replace("&rdquo;", "”")
    text = text.replace("&middot;", "·")
    parts = re.split(r"(<strong>.*?</strong>|<em>.*?</em>|<code>.*?</code>)", text)
    for p in parts:
        if not p:
            continue
        if p.startswith("<strong>"):
            r = par.add_run(strip_tags(p)); r.bold = True
        elif p.startswith("<em>"):
            r = par.add_run(strip_tags(p)); r.italic = True
        elif p.startswith("<code>"):
            r = par.add_run(strip_tags(p)); r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            par.add_run(strip_tags(p))


def set_table_borders(tbl):
    """Horizontal rules only: top and bottom of the table, and under the header."""
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    spec = {"top": ("single", 12), "bottom": ("single", 12),
            "left": ("none", 0), "right": ("none", 0),
            "insideH": ("none", 0), "insideV": ("none", 0)}
    for edge, (val, sz) in spec.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "17222E")
        borders.append(el)
    tblPr.append(borders)


def header_rule(row):
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        bo = OxmlElement("w:bottom")
        bo.set(qn("w:val"), "single"); bo.set(qn("w:sz"), "8")
        bo.set(qn("w:color"), "17222E")
        b.append(bo); tcPr.append(b)


def render_docx(bl):
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    st = d.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.25

    t = d.add_heading(C.TITLE, level=0)
    for r in t.runs:
        r.font.color.rgb = RGBColor(0x17, 0x22, 0x2E)
    p = d.add_paragraph(); r = p.add_run(C.SUBTITLE); r.italic = True
    r.font.size = Pt(11)
    p = d.add_paragraph()
    r = p.add_run(f"{C.AUTHOR}  ·  Roll no. {C.ROLL}\n"); r.bold = True
    r2 = p.add_run(f"{C.AFFIL}\nCapstone Project 10")
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(0x73, 0x81, 0x8D)

    for blk in bl:
        k = blk[0]
        if k == "h2":
            d.add_heading(strip_tags(blk[1]), level=1)
        elif k == "h3":
            d.add_heading(strip_tags(blk[1]), level=2)
        elif k == "p":
            rich_runs(d.add_paragraph(), blk[1])
        elif k == "note":
            par = d.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            rich_runs(par, blk[1])
            for r in par.runs:
                r.font.size = Pt(9.8)
                r.font.color.rgb = RGBColor(0x9C, 0x3D, 0x1E)
        elif k == "bullets":
            for item in blk[1]:
                par = d.add_paragraph(style="List Bullet")
                rich_runs(par, item)
        elif k == "pre":
            if blk[1].strip():
                for line in blk[1].split("\n"):
                    par = d.add_paragraph()
                    par.paragraph_format.space_after = Pt(0)
                    par.paragraph_format.line_spacing = 1.0
                    r = par.add_run(line)
                    r.font.name = "Consolas"; r.font.size = Pt(7.2)
        elif k == "fig":
            d.add_picture(str(ROOT / blk[1]), width=Cm(16.0))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = d.add_paragraph()
            r = cap.add_run(strip_tags(blk[2]))
            r.font.size = Pt(9); r.italic = True
            r.font.color.rgb = RGBColor(0x44, 0x52, 0x5F)
        elif k == "table":
            _, caption, hdr, rows = blk
            cp = d.add_paragraph()
            r = cp.add_run(strip_tags(caption))
            r.font.size = Pt(9); r.italic = True
            r.font.color.rgb = RGBColor(0x44, 0x52, 0x5F)
            tbl = d.add_table(rows=1, cols=len(hdr))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(tbl)
            hc = tbl.rows[0].cells
            for i, h in enumerate(hdr):
                hc[i].text = ""
                par = hc[i].paragraphs[0]
                run = par.add_run(strip_tags(h))
                run.bold = True; run.font.size = Pt(8.2)
                par.paragraph_format.space_after = Pt(3)
                if i:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            header_rule(tbl.rows[0])
            for row in rows:
                cells = tbl.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = ""
                    par = cells[i].paragraphs[0]
                    run = par.add_run(strip_tags(str(val)))
                    run.font.size = Pt(8.2)
                    par.paragraph_format.space_after = Pt(2)
                    if i:
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    make_excerpt()
    import importlib
    importlib.reload(C)
    bl = C.blocks()
    h = render_html(bl)
    dd = render_docx(bl)
    print(f"blocks: {len(bl)}")
    print(f"wrote {h}")
    print(f"wrote {dd}")
